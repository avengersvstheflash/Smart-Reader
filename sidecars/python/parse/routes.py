import io
from fastapi import APIRouter, File, UploadFile, HTTPException
from docx import Document
from striprtf.striprtf import rtf_to_text

import base64
from pydantic import BaseModel
from .structure import detect_structure

router = APIRouter(prefix="/v1/parse", tags=["parse"])

class StructureRequest(BaseModel):
    pdf_base64: str

from fastapi.concurrency import run_in_threadpool

def _parse_docx_sync(file_bytes):
    doc = Document(io.BytesIO(file_bytes))
    title = doc.core_properties.title or ""
    author = doc.core_properties.author or "Unknown Author"
    
    blocks = []
    for element in doc.element.body:
        if element.tag.endswith('p'):
            for p in doc.paragraphs:
                if p._p == element:
                    text = p.text.strip()
                    if not text:
                        break
                    style_name = p.style.name if p.style else ""
                    if style_name.startswith("Heading"):
                        try:
                            level = int(style_name.replace("Heading", "").strip())
                            blocks.append({"kind": "heading", "level": level, "text": text})
                        except:
                            blocks.append({"kind": "paragraph", "text": text})
                    else:
                        blocks.append({"kind": "paragraph", "text": text})
                    break
        elif element.tag.endswith('tbl'):
            for t in doc.tables:
                if t._tbl == element:
                    rows = []
                    for row in t.rows:
                        rows.append([cell.text.strip() for cell in row.cells])
                    blocks.append({"kind": "table", "rows": rows})
                    break
    
    if not title:
        for b in blocks:
            if b["kind"] == "heading":
                title = b["text"]
                break

    return {"title": title, "author": author, "blocks": blocks}

@router.post("/docx")
async def parse_docx_endpoint(file: UploadFile = File(...)):
    file_bytes = await file.read()
    print(f"[Routing: Parse] Endpoint: /docx | Bytes: {len(file_bytes)}")
    try:
        return await run_in_threadpool(_parse_docx_sync, file_bytes)
    except Exception as e:
        raise HTTPException(status_code=400, detail={"status": "error", "code": "INVALID_DOCX", "message": str(e)})

def _parse_rtf_sync(file_bytes):
    text = rtf_to_text(file_bytes.decode('utf-8', errors='replace'))
    paragraphs = [p.strip() for p in text.split('\n\n') if p.strip()]
    
    blocks = [{"kind": "paragraph", "text": p} for p in paragraphs]
    
    title = ""
    author = "Unknown Author"
    
    # Super basic info extraction
    import re
    info_match = re.search(r'{\\info.*?}', file_bytes.decode('utf-8', errors='replace'), re.DOTALL)
    if info_match:
        info_str = info_match.group(0)
        t_match = re.search(r'{\\title(.*?)}', info_str)
        if t_match: title = t_match.group(1).strip()
        a_match = re.search(r'{\\author(.*?)}', info_str)
        if a_match: author = a_match.group(1).strip()
        
    if not title and paragraphs:
        title = paragraphs[0]
        
    return {"title": title, "author": author, "blocks": blocks}

@router.post("/rtf")
async def parse_rtf_endpoint(file: UploadFile = File(...)):
    file_bytes = await file.read()
    print(f"[Routing: Parse] Endpoint: /rtf | Bytes: {len(file_bytes)}")
    try:
        return await run_in_threadpool(_parse_rtf_sync, file_bytes)
    except Exception as e:
        raise HTTPException(status_code=400, detail={"status": "error", "code": "INVALID_RTF", "message": str(e)})


@router.post("/structure")
def parse_structure_endpoint(payload: StructureRequest):
    print(f"[Routing: Parse] Endpoint: /structure | Base64Length: {len(payload.pdf_base64)}")
    try:
        pdf_bytes = base64.b64decode(payload.pdf_base64)
    except Exception as e:
        return {
            "has_embedded_toc": False,
            "toc": [],
            "heading_candidates": [],
            "front_matter_page_range": None,
            "back_matter_page_range": None,
            "estimated_chapter_count": 0,
            "method": "none",
            "warnings": [f"invalid_base64_payload: {str(e)}"],
        }
    return detect_structure(pdf_bytes)

